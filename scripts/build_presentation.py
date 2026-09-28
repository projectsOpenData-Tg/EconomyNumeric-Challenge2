"""Génère le rapport PowerPoint du Défi 2 — Économie numérique, écrit pour le décideur.

Plan : `tmp/ppt-plan.md`. Processus repris du défi 1 (`EconomyNumeric-Challenge-1/scripts/build_presentation.py`) :
un plan, un script, des contrôles. La relecture « décideur » qui avait fait réécrire les dix diapositives du défi 1 est
appliquée ici dès la première version : chaque diapositive dit ce qu'il faut décider, pas ce qui a été analysé.

⚠️ **Aucun chiffre n'est retapé.** Toutes les valeurs sont lues dans `data/analysis/` par `dashboard/donnees.py`, comme
le fait le tableau de bord : deux endroits qui portent la même grandeur finissent par diverger. Les constats qui nomment
des territoires (« toutes hors du Maritime », « 14 des 22 communes ») sont dérivés des tables, pour rester justes si le
pipeline est rejoué. Les couleurs viennent de `dashboard/theme.py` : le rapport et le dashboard ont la même identité.

⚠️ **Langage des décideurs** (`workspace/_GARDRAILS_DASHBOARD.txt`, consignes du 28/09/2026) : aucun code interne, aucun
nom de fichier, aucun sigle, aucun nom de source sur les diapositives. La diapositive 2 est la seule à parler de méthode,
en clair ; le détail reste sur la page « Sources et méthode » du dashboard.

Contraintes techniques héritées du défi 1 : graphiques en matplotlib (l'export d'images Plotly exige Chrome) ; armoiries
rasterisées une fois (`python-pptx` ne lit pas le SVG).

Usage  : `python3 scripts/build_presentation.py`
Sortie : `Togo-Economie-Numerique-Defi2.pptx`, à la racine du projet.
"""

from __future__ import annotations

import datetime
import html
import logging
import subprocess
import sys
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import geopandas as gpd  # noqa: E402
import matplotlib.pyplot as plt  # noqa: E402
import numpy as np  # noqa: E402
import pandas as pd  # noqa: E402
from matplotlib import font_manager  # noqa: E402
from matplotlib.lines import Line2D  # noqa: E402
from matplotlib.patches import Patch  # noqa: E402
from PIL import Image, ImageFont  # noqa: E402
from pptx import Presentation  # noqa: E402
from pptx.dml.color import RGBColor  # noqa: E402
from pptx.enum.shapes import MSO_SHAPE  # noqa: E402
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN  # noqa: E402
from pptx.util import Emu, Inches, Pt  # noqa: E402

RACINE = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(RACINE / "dashboard"))

import i18n  # noqa: E402
from donnees import GEO, REGIONS, chiffres_nationaux, communes, fr, lire, prefectures  # noqa: E402
from theme import (BLEUS, COULEUR_REGION, DISCRET, ENCRE, ENCRE2, FILET, OCRE, PRIORITE, STATUT_O4_05,  # noqa: E402
                   TEXTE_CATEGORIELLE)

# Hors d'une session Streamlit, les lectures en cache du dashboard signalent l'absence de serveur à chaque appel.
for _nom in list(logging.root.manager.loggerDict):
    if _nom.startswith("streamlit"):
        logging.getLogger(_nom).setLevel(logging.ERROR)

ASSETS = RACINE / "tmp" / "ppt_assets"
ASSETS.mkdir(parents=True, exist_ok=True)
SORTIE = RACINE / "Togo-Economie-Numerique-Defi2.pptx"

# --------------------------------------------------------------------------------------------------------------------
# Identité : textes et adresses
# --------------------------------------------------------------------------------------------------------------------

#: ⚠️ À vérifier avant l'envoi : le nom paraît sur la couverture (note du fichier `.env` : « Edmund »).
PRESENTE_PAR = "Edmund"

#: ⚠️ À vérifier avant l'envoi : application Heroku relevée dans le fichier `.env` ; aucun déploiement n'est visible
#: dans le dépôt. Laissée vide, la ligne n'est pas rendue — jamais de texte factice.
URL_DASHBOARD = "https://togo-econum-finance-defi2-5d98fee2ba7c.herokuapp.com"

#: Mention de source des gardrails, section « Sources » : générique sur les pages de décision.
LIGNE_SOURCE = ("Données : portail national de données ouvertes du Togo, recensement de 2022 et sources "
                "institutionnelles complémentaires.")
PIED_IDENTITE = "Togo AI Lab — Data Challenge | Économie numérique — Défi 2"
BANDEAU = "Rapport de synthèse · Économie numérique — Défi 2"

# Textes de la barre du haut du dashboard, lus dans son dictionnaire plutôt que recopiés.
MINISTERE = i18n._T["topbar.ministry"]["fr"]
MARQUE = html.unescape(i18n._T["topbar.brand"]["fr"])
PAGE_METHODE = i18n._T["page.methodologie"]["fr"]

# Règles du projet (02 et 08), pas des résultats : elles servent à lire les chiffres, et sont contrôlées plus bas.
SEUIL_PRIORITE_HAUTE = 70     # score de 70 ou plus : priorité haute (08, section 3)
SEUIL_USAGE = 40              # part de la population utilisant Internet : sortie du « déficit d'usage » (O1-01)
SEUIL_INVESTISSEMENT = 15     # part du chiffre d'affaires investie : seuil de sous-investissement (O2-04)

# --------------------------------------------------------------------------------------------------------------------
# Palette : celle du dashboard
# --------------------------------------------------------------------------------------------------------------------

NAVY = PRIORITE["haute"]                 # #0d366b : titres de carte, priorité haute, phrases de lecture
BLANC = "#FFFFFF"
JAUNE = "#fce588"                        # jaune clair des titres de filtres, sur fond sombre
OR = "#ffce00"                           # bordure de la carte du logo, dans la barre du haut
OCRE_TEXTE = TEXTE_CATEGORIELLE[OCRE]    # #ab6300 : ocre lisible sur fond blanc (surtitres)
GRIS_POINT = "#b9b6ad"
BLEU_CLAIR_TEXTE = "#cde2fb"
VERT_SERIE, ROUGE_SERIE, ORANGE_SERIE = "#1baf7a", "#e34948", "#eb6834"

#: Tons des étiquettes et des encadrés du dashboard (theme.py : .etiquette, .constat, .limite, thèmes des recommandations).
TONS = {
    "bleu":   dict(bg="#e8f0fb", bord="#c7d8f0", encre=NAVY),
    "navy":   dict(bg="#e3eefb", bord="#c7d8f0", encre=NAVY),
    "ambre":  dict(bg="#fdf0d2", bord="#ecd28f", encre="#6b4700"),
    "limite": dict(bg="#fdf3d7", bord="#ecd28f", encre="#6b4700"),
    "rouge":  dict(bg="#fbe3e2", bord="#f1c4c2", encre="#8a1c1b"),
    "vert":   dict(bg="#e2f4ec", bord="#bfe3d2", encre="#11613f"),
    "violet": dict(bg="#ece9fb", bord="#d5cff3", encre="#4a3aa7"),
    "orange": dict(bg="#fde8dd", bord="#f6cdb7", encre="#b4451a"),
    "gris":   dict(bg="#efece4", bord=FILET, encre=ENCRE2),
}
ETIQUETTE = {"alerte": "ambre", "ok": "navy", "critique": "rouge", "neutre": "gris"}

#: Les trois dimensions du classement, dans les couleurs de leurs séries (figures du 08), en pas lisibles sur blanc.
DIM_TEXTE = {"acces": TEXTE_CATEGORIELLE["#eb6834"], "maillage": "#11613f", "couverture": "#4a3aa7"}
DIM_POINT = {"acces": "#eb6834", "maillage": "#1baf7a", "couverture": "#4a3aa7"}

POLICE, POLICE_TITRE = "Calibri", "Cambria"
LARGEUR, HAUTEUR = Emu(12192000), Emu(6858000)   # 16:9 exact (défi 1 : Inches(13.333) laissait 305 EMU de moins)
SW, SH = 13.333, 7.5
M = 0.55                                        # marges gauche et droite
CW = SW - 2 * M                                 # largeur utile : 12,233
DROITE = M + CW
CORPS_Y, CORPS_Y_SOUS = 1.60, 1.96             # haut du contenu, sans et avec sous-titre
PIED_Y = 7.05

AVERTISSEMENTS: list[str] = []                  # textes réduits pour tenir dans leur cadre, signalés en fin d'exécution

# --------------------------------------------------------------------------------------------------------------------
# Données — lues une fois, jamais retapées
# --------------------------------------------------------------------------------------------------------------------

N = chiffres_nationaux()
P = prefectures()
C = communes()
SR = lire("10_recommandations", "synthese_recommandations").set_index("id")

# Objectif 1
D1 = lire("05_eda", "s4_usage_internet_d1").set_index("annee")
USAGE = lire("07_indicateurs", "o1_02_usage").set_index("annee")
ACCES = lire("06_spatial", "s6_usage_internet_regions").set_index("nom")
SCEN = lire("10_recommandations", "r8_scenarios_synthese")
A_U = int(USAGE.index.max())
MULT_2017 = D1.loc[A_U, "pct_population"] / D1.loc[2017, "pct_population"]
ANNEES_ACCEL = [int(a) for a in USAGE.index if USAGE.loc[a, "classe"] == "accélération"]
VAR_PIC = USAGE.loc[max(ANNEES_ACCEL), "variation_points"]
# Dernières années consécutives sous 2 points gagnés : « moins de 2 points par an depuis … »
_depuis = A_U
while _depuis - 1 in USAGE.index and USAGE.loc[_depuis - 1, "variation_points"] < 2:
    _depuis -= 1
DEPUIS_LENT = _depuis
TENDANCIEL = SCEN[SCEN.scenario.str.startswith("tendanciel")].iloc[0]
ACCELERE = SCEN[SCEN.scenario.str.startswith("accéléré")].iloc[0]
ACCES_GL, ACCES_SAV = ACCES.loc["Grand Lomé", "estimation_pct_2021_22"], ACCES.loc["Savanes", "estimation_pct_2021_22"]
ACCES_PLA = ACCES.loc["Plateaux", "estimation_pct_2021_22"]

# Objectif 2
INV = lire("07_indicateurs", "o2_04_investissement").set_index("annee")
SITES = lire("07_indicateurs", "o2_08_sites_radio").set_index("annee")
R8 = lire("10_recommandations", "r8_investissement").iloc[0]
R10 = lire("10_recommandations", "r10_sites_radio").iloc[0]
R6 = lire("10_recommandations", "r6_cout_data").iloc[0]
A_I0, A_I = int(INV.index.min()), int(INV.index.max())
SITES_S = SITES.ajouts_nets_total.dropna()
REPERE_SITES = int(round(R10.seuil_declare_proche_de_zero))   # 49,5 → 50, comme le dashboard (P26 du 10)

# Objectif 3
PARTS = lire("05_eda", "s2_parts_unites_regionales").set_index("nom")
STRATES = lire("07_indicateurs", "o3_02_strates")
STRATES = STRATES[STRATES.strates.str.contains("50")].set_index("strate")
BF = lire("07_indicateurs", "o3_04_bceao_findex").set_index("mesure").valeur
FRAIS = lire("07_indicateurs", "o3_06_frais")
FRAIS = FRAIS[FRAIS.montant_fcfa.notna()].set_index("montant_fcfa")
R5 = lire("10_recommandations", "r5_frais_mobile_money").iloc[0]
DETENTION = BF[[i for i in BF.index if i.startswith("Détention")][0]]
# Parts par milieu, sommées sur les communes (même seuil de 50 % d'urbains que le 07) : population, agences, mobile money.
PAR_MILIEU = C.groupby("strate_50")[["pop_totale", "n_formels", "n_mm"]].sum()
PAR_MILIEU = 100 * PAR_MILIEU / PAR_MILIEU.sum()
assert (PAR_MILIEU.n_formels - STRATES.part_points_pct).abs().max() < 0.05, "parts d'agences : écart avec le 07"
COMMUNES_MM = int((C.n_mm > 0).sum())

# Objectif 4
MMU = C[C.statut_O4_05 == "mobile money uniquement"]
SUPPL = C[C.classe_O4_03 == "suppléance quasi totale"]
CRIT = C[C.cellule_critique]
PART_SUPPL = 100 * SUPPL.pop_totale.sum() / C.pop_totale.sum()

# Priorités (objectif 5)
P1 = P[P.priorite == "haute"]
NC = P[P.priorite == "non classée"]
CLASSEES = P[P.priorite != "non classée"]
PART_P1 = 100 * P1.pop_totale.sum() / CLASSEES.pop_totale.sum()
HORS_MARITIME = not (set(P1.unite_regionale) & {"Grand Lomé", "Maritime hors Grand Lomé"})
assert P1.score.min() >= SEUIL_PRIORITE_HAUTE > P[P.priorite == "moyenne"].score.max(), "seuil de la priorité haute"
FICHES = lire("09_diagnostic", "fiches_prefectures").set_index("nom")
FACT = lire("09_diagnostic", "facteurs_repetition").set_index("colonne").loc["part_points_mm_plus_10km_guichet_pct"]

# Recommandations
R1 = lire("10_recommandations", "r1_communes_mobile_money_uniquement").sort_values("ordre_action")
R2 = lire("10_recommandations", "r2_prefectures_points_formels")
R3 = lire("10_recommandations", "r3_maillage_mobile_money")
R4A = lire("10_recommandations", "r4_communes_a_mesurer")
R4B = lire("10_recommandations", "r4_couverture")
R4B_P1 = R4B[R4B.classe_08 == "priorité 1"]
R9 = lire("10_recommandations", "r9_fibre_non_raccordees")
R7 = lire("10_recommandations", "r7_competences")
R3P, R3C = R3[R3.maille == "préfecture"], R3[R3.maille == "commune"]
# Points mobile money qui franchissent un seuil du 02 ; la médiane des préfectures reste un ordre de grandeur (P22).
MM_SEUIL = int(R3[~R3.cible.str.startswith("médiane")].points_a_ajouter.sum())
# Préfectures dont toutes les communes sont à mesurer.
_a_mesurer = set(R4A.code)
MESUREES_ENTIERES = [p for p, g in C.groupby("prefecture") if set(g.code) <= _a_mesurer]
MESUREES_ENTIERES = sorted(MESUREES_ENTIERES, key=lambda p: -C[C.prefecture == p].pop_totale.sum())
# Communes sans agence situées dans des préfectures qui ne sont pas en priorité haute (ni non classées).
MMU_AILLEURS = R1[R1.classe_prefecture_08.isin(["priorité 2", "priorité 3"])]
TOP3_POP = P.nlargest(3, "pop_totale")


# --------------------------------------------------------------------------------------------------------------------
# Mise en forme des nombres (espaces insécables : un nombre ne se coupe jamais en fin de ligne)
# --------------------------------------------------------------------------------------------------------------------

def nb(x, d=0) -> str:
    return fr(x, d).replace(" ", " ")


def pct(x, d=1) -> str:
    return f"{nb(x, d)} %"


def millions(x, d=2) -> str:
    return f"{nb(x / 1e6, d)} million{'s' if x >= 2e6 else ''}"


def et(noms: list[str]) -> str:
    noms = list(noms)
    return noms[0] if len(noms) == 1 else ", ".join(noms[:-1]) + " et " + noms[-1]


def rgb(couleur: str) -> RGBColor:
    c = couleur.lstrip("#")
    return RGBColor(int(c[0:2], 16), int(c[2:4], 16), int(c[4:6], 16))


# --------------------------------------------------------------------------------------------------------------------
# Mesure du texte — pour ajuster la taille avant d'écrire, pas après
# --------------------------------------------------------------------------------------------------------------------

#: Métriques d'Arial (Liberation Sans), volontairement pessimistes : Calibri est plus étroite d'environ 10 %, Cambria
#: a une largeur voisine. Un texte qui tient ici tient dans PowerPoint.
_POLICES_MESURE = {False: "/usr/share/fonts/liberation-sans/LiberationSans-Regular.ttf",
                   True: "/usr/share/fonts/liberation-sans/LiberationSans-Bold.ttf"}
_cache_polices: dict = {}


def largeur_texte(texte: str, taille: float, gras: bool = False) -> float:
    """Largeur d'un texte, en pouces."""
    if gras not in _cache_polices:
        try:
            _cache_polices[gras] = ImageFont.truetype(_POLICES_MESURE[gras], 100)
        except OSError:
            _cache_polices[gras] = None
    police = _cache_polices[gras]
    if police is None:
        return len(texte) * 0.55 * taille / 72
    return police.getlength(texte) / 100 * taille / 72


def nombre_lignes(texte: str, taille: float, gras: bool, largeur: float) -> int:
    lignes = 0
    for morceau in texte.split("\n"):
        courante, n = "", 1
        for mot in morceau.split(" "):
            essai = f"{courante} {mot}" if courante else mot
            if courante and largeur_texte(essai, taille, gras) > largeur:
                n, courante = n + 1, mot
            else:
                courante = essai
        lignes += n
    return lignes


def taille_ajustee(texte: str, largeur: float, tailles=(26, 24, 22, 20), gras=True) -> float:
    """La plus grande taille à laquelle le texte tient sur une seule ligne."""
    for taille in tailles:
        if largeur_texte(texte, taille, gras) <= largeur:
            return taille
    return tailles[-1]


# --------------------------------------------------------------------------------------------------------------------
# Primitives de composition — toutes les coordonnées sont en pouces
# --------------------------------------------------------------------------------------------------------------------

def nouvelle_diapositive(presentation):
    return presentation.slides.add_slide(presentation.slide_layouts[6])


def rect(d, x, y, w, h, *, fond=None, rayon=0.0, bord=None, epaisseur=0.75):
    forme = d.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if rayon else MSO_SHAPE.RECTANGLE,
                               Inches(x), Inches(y), Inches(w), Inches(h))
    if rayon:
        forme.adjustments[0] = min(0.5, rayon / min(w, h))
    if fond:
        forme.fill.solid()
        forme.fill.fore_color.rgb = rgb(fond)
    else:
        forme.fill.background()
    if bord:
        forme.line.color.rgb = rgb(bord)
        forme.line.width = Pt(epaisseur)
    else:
        forme.line.fill.background()
    forme.shadow.inherit = False
    return forme


def _sans_marges(cadre) -> None:
    cadre.margin_left = cadre.margin_right = cadre.margin_top = cadre.margin_bottom = 0


def _run(paragraphe, texte, *, taille=12, gras=False, couleur=ENCRE, italique=False, police=POLICE, lien=None):
    run = paragraphe.add_run()
    run.text = texte
    run.font.size = Pt(taille)
    run.font.bold = gras
    run.font.italic = italique
    run.font.name = police
    run.font.color.rgb = rgb(couleur)
    if lien:
        run.hyperlink.address = lien
        run.font.underline = True
        run.font.color.rgb = rgb(couleur)
    return run


_STYLE = ("taille", "gras", "couleur", "italique", "police", "lien")


def _runs(ligne: dict) -> list[dict]:
    base = {k: ligne[k] for k in _STYLE if k in ligne}
    if "runs" in ligne:
        return [dict(base, **r) for r in ligne["runs"]]
    return [dict(base, texte=ligne["texte"])]


def _hauteur(lignes: list[dict], largeur: float, facteur: float = 1.0) -> float:
    total = 0.0
    for ligne in lignes:
        runs = _runs(ligne)
        taille = max(r.get("taille", 12) for r in runs) * facteur
        gras = any(r.get("gras") for r in runs)
        n = nombre_lignes("".join(r["texte"] for r in runs), taille, gras, largeur)
        total += n * taille * 1.2 * ligne.get("interligne", 1.0) / 72
        total += (ligne.get("avant", 0) + ligne.get("apres", 0)) / 72
    return total


def bloc(d, x, y, w, h, lignes, *, align=PP_ALIGN.LEFT, ancre=MSO_ANCHOR.TOP, ajuster=True):
    """Une zone de texte. Chaque ligne est un paragraphe : `texte` ou `runs`, et son style.

    Interligne et espacements sont **explicites, en points** : laissés implicites, chaque moteur de rendu les recalcule
    (leçon du défi 1). Si le texte ne tient pas, la taille est réduite par pas de 2,5 %, jusqu'à 80 % au plus, et la
    réduction est signalée en fin d'exécution : un texte qui déborde n'est jamais laissé tel quel.
    """
    lignes = [dict(l) if isinstance(l, dict) else dict(texte=l) for l in lignes]
    facteur = 1.0
    if ajuster:
        while _hauteur(lignes, w, facteur) > h + 0.01 and facteur > 0.80:
            facteur -= 0.025
        if facteur < 1.0:
            apercu = "".join(r["texte"] for r in _runs(lignes[0]))[:50]
            AVERTISSEMENTS.append(f"réduit à {facteur:.0%} : « {apercu}… »")
    boite = d.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    cadre = boite.text_frame
    cadre.word_wrap = True
    cadre.auto_size = MSO_AUTO_SIZE.NONE
    cadre.vertical_anchor = ancre
    _sans_marges(cadre)
    for i, ligne in enumerate(lignes):
        p = cadre.paragraphs[0] if i == 0 else cadre.add_paragraph()
        p.alignment = ligne.get("align", align)
        p.line_spacing = ligne.get("interligne", 1.0)
        p.space_before = Pt(ligne.get("avant", 0))
        p.space_after = Pt(ligne.get("apres", 0))
        for r in _runs(ligne):
            style = {k: r[k] for k in _STYLE if k in r}
            style["taille"] = round(style.get("taille", 12) * facteur * 2) / 2
            _run(p, r["texte"], **style)
    return boite


def T(texte, **style) -> dict:
    return dict(texte=texte, **style)


def R(region: str, **style) -> dict:
    """Un nom de région écrit dans sa couleur, comme dans les tableaux et les constats du dashboard."""
    return dict(texte=region, couleur=COULEUR_REGION[region], gras=True, **style)


def surtitre(d, x, y, w, texte, couleur=OCRE_TEXTE):
    bloc(d, x, y, w, 0.24, [T(texte.upper(), taille=10, gras=True, couleur=couleur)])


def pastille(d, x, y, texte, ton, *, taille=9, h=0.24, droite=False) -> float:
    """Étiquette arrondie, comme `.etiquette` du dashboard. Avec `droite`, x est le bord droit. Renvoie la largeur."""
    w = largeur_texte(texte, taille, True) + 0.26
    if droite:
        x -= w
    tonalite = TONS[ETIQUETTE.get(ton, ton)]
    forme = rect(d, x, y, w, h, fond=tonalite["bg"], rayon=h / 2)
    cadre = forme.text_frame
    cadre.word_wrap = False
    _sans_marges(cadre)
    cadre.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = cadre.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    _run(p, texte, taille=taille, gras=True, couleur=tonalite["encre"])
    return w


def carre_numero(d, x, y, cote, texte, couleur, *, taille=13):
    """Le motif du rapport : une pastille carrée arrondie, pleine, chiffre blanc (pastilles de thème du dashboard)."""
    forme = rect(d, x, y, cote, cote, fond=couleur, rayon=cote * 0.28)
    cadre = forme.text_frame
    _sans_marges(cadre)
    cadre.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = cadre.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    _run(p, texte, taille=taille, gras=True, couleur=BLANC, police=POLICE_TITRE)
    return forme


def carte(d, x, y, w, h, *, fond=BLANC, bord=FILET):
    """Carte blanche, bordure fine, coins arrondis : le conteneur du dashboard."""
    return rect(d, x, y, w, h, fond=fond, rayon=0.12, bord=bord)


def cartouche(d, x, y, w, h, *, surtitre_txt, valeur, unite="", phrase, note="", etiquette=None, ton="alerte",
              couleur_valeur=ENCRE):
    """Carte de chiffre clé, sur le modèle de `.kpi` du dashboard : surtitre, étiquette, valeur, phrase, note."""
    carte(d, x, y, w, h)
    pad = 0.18
    interieur = w - 2 * pad
    largeur_etiquette = pastille(d, x + w - pad, y + 0.13, etiquette, ton, taille=8.5, h=0.22, droite=True) + 0.08 \
        if etiquette else 0
    bloc(d, x + pad, y + 0.15, interieur - largeur_etiquette, 0.22,
         [T(surtitre_txt.upper(), taille=8.5, gras=True, couleur=NAVY)])
    runs = [T(valeur, taille=25, gras=True, couleur=couleur_valeur, police=POLICE_TITRE)]
    if unite:
        runs.append(T(f" {unite}", taille=12, gras=True, couleur=couleur_valeur))
    bloc(d, x + pad, y + 0.40, interieur, 0.46, [dict(runs=runs)], ajuster=False)
    # La phrase prend une ou deux lignes : la note se place dessous, jamais sur elle.
    hauteur_phrase = nombre_lignes(phrase, 10.5, True, interieur) * 10.5 * 1.2 / 72 + 0.02
    bloc(d, x + pad, y + 0.88, interieur, hauteur_phrase, [T(phrase, taille=10.5, gras=True, couleur=ENCRE)])
    if note:
        y_note = y + 0.88 + hauteur_phrase + 0.05
        bloc(d, x + pad, y_note, interieur, y + h - 0.08 - y_note, [T(note, taille=9, couleur=ENCRE2)])


def banniere(d, x, y, w, h, etiquette, texte, style="sombre", *, taille=11.5):
    """Bandeau de message : `sombre` (le message à retenir), `constat`, `limite` ou `vert`."""
    if style == "sombre":
        rect(d, x, y, w, h, fond=NAVY, rayon=0.10)
        encre, encre_etiquette, gras = BLANC, JAUNE, True
    else:
        ton = {"constat": "bleu", "limite": "limite", "vert": "vert"}[style]
        rect(d, x, y, w, h, fond=TONS[ton]["bg"], rayon=0.10, bord=TONS[ton]["bord"])
        encre, encre_etiquette, gras = ENCRE, TONS[ton]["encre"], False
    runs = []
    if etiquette:
        runs.append(T(f"{etiquette.upper()}  ·  ", taille=taille - 1, gras=True, couleur=encre_etiquette))
    runs += [dict(r, **{"couleur": r.get("couleur", encre)}) for r in (texte if isinstance(texte, list) else [T(texte)])]
    runs = [dict(r, taille=r.get("taille", taille), gras=r.get("gras", gras)) for r in runs]
    bloc(d, x + 0.26, y + 0.06, w - 0.52, h - 0.12, [dict(runs=runs, interligne=1.05)], ancre=MSO_ANCHOR.MIDDLE)


def encadre(d, x, y, w, h, ton, runs, *, taille=11.5, gras=True):
    """Encadré teinté (constat, enjeu) : texte en gras, chiffres et régions en couleur."""
    rect(d, x, y, w, h, fond=TONS[ton]["bg"], rayon=0.12, bord=TONS[ton]["bord"])
    runs = [dict(r, taille=r.get("taille", taille), gras=r.get("gras", gras)) for r in runs]
    bloc(d, x + 0.22, y + 0.12, w - 0.44, h - 0.24, [dict(runs=runs, interligne=1.12)], ancre=MSO_ANCHOR.MIDDLE)


def panneau(d, x, y, w, h, titre, *, sous_titre=None) -> float:
    """Carte blanche titrée, pour un graphique ou une carte. Renvoie l'ordonnée sous le titre."""
    carte(d, x, y, w, h)
    bloc(d, x + 0.22, y + 0.14, w - 0.44, 0.28, [T(titre, taille=12.5, gras=True, couleur=ENCRE, police=POLICE_TITRE)])
    if sous_titre:
        bloc(d, x + 0.22, y + 0.42, w - 0.44, 0.22, [T(sous_titre, taille=9, couleur=DISCRET)])
        return y + 0.66
    return y + 0.46


def lien(d, x, y, w, texte, url, *, taille=11, couleur=NAVY):
    boite = d.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(0.28))
    cadre = boite.text_frame
    cadre.word_wrap = False
    _sans_marges(cadre)
    _run(cadre.paragraphs[0], texte, taille=taille, gras=True, couleur=couleur, lien=url)


def image_ajustee(d, chemin: Path, x, y, largeur_max, hauteur_max, *, centre=True):
    with Image.open(chemin) as image:
        largeur, hauteur = image.size
    ratio = min(largeur_max / largeur, hauteur_max / hauteur)
    w, h = largeur * ratio, hauteur * ratio
    if centre:
        x, y = x + (largeur_max - w) / 2, y + (hauteur_max - h) / 2
    d.shapes.add_picture(str(chemin), Inches(x), Inches(y), width=Inches(w), height=Inches(h))


def tableau_texte(d, x, y, w, colonnes, lignes, *, taille=10.5, h_ligne=0.30, h_entete=0.28) -> float:
    """Tableau dessiné en zones de texte (plus léger qu'un tableau PowerPoint). `colonnes` : (titre, largeur, alignement).
    Une cellule est un texte, ou un tuple (texte, couleur, gras). Renvoie l'ordonnée sous la dernière ligne."""
    rect(d, x, y, w, h_entete, fond=TONS["gris"]["bg"], rayon=0.06)
    curseur = x + 0.14
    for titre, largeur, align in colonnes:
        bloc(d, curseur, y, largeur - 0.08, h_entete, [T(titre, taille=8.5, gras=True, couleur=ENCRE2)], align=align,
             ancre=MSO_ANCHOR.MIDDLE)
        curseur += largeur
    ordonnee = y + h_entete
    for i, ligne in enumerate(lignes):
        curseur = x + 0.14
        for (_, largeur, align), cellule in zip(colonnes, ligne):
            texte, couleur, gras = cellule if isinstance(cellule, tuple) else (cellule, ENCRE, False)
            bloc(d, curseur, ordonnee + 0.02, largeur - 0.08, h_ligne - 0.04,
                 [T(texte, taille=taille, couleur=couleur, gras=gras, interligne=0.95)], align=align,
                 ancre=MSO_ANCHOR.MIDDLE)
            curseur += largeur
        ordonnee += h_ligne
        if i < len(lignes) - 1:
            rect(d, x + 0.10, ordonnee - 0.004, w - 0.20, 0.008, fond=FILET)
    return ordonnee


def en_tete(d, numero: int, surtitre_txt: str, titre: str, *, sous_titre: str | None = None) -> float:
    """Armoiries et nom du projet (la barre du haut du dashboard), surtitre, titre, numéro. Renvoie le haut du contenu."""
    image_ajustee(d, ARMOIRIES, M, 0.17, 0.44, 0.40, centre=False)
    bloc(d, M + 0.56, 0.20, 6.2, 0.32, [T(MARQUE.upper(), taille=10, gras=True, couleur=NAVY)], ancre=MSO_ANCHOR.MIDDLE)
    bloc(d, DROITE - 5.6, 0.20, 5.6, 0.32, [T(BANDEAU, taille=9.5, couleur=DISCRET)], align=PP_ALIGN.RIGHT,
         ancre=MSO_ANCHOR.MIDDLE)
    surtitre(d, M, 0.70, 9.0, surtitre_txt)
    bloc(d, DROITE - 0.8, 0.66, 0.8, 0.30, [T(str(numero), taille=14, gras=True, couleur=GRIS_POINT)],
         align=PP_ALIGN.RIGHT)
    bloc(d, M, 0.95, CW, 0.52, [T(titre, taille=taille_ajustee(titre, CW), gras=True, couleur=ENCRE,
                                   police=POLICE_TITRE)])
    if sous_titre:
        bloc(d, M, 1.50, CW, 0.28, [T(sous_titre, taille=12, couleur=DISCRET)])
        return CORPS_Y_SOUS
    return CORPS_Y


def pied(d, texte: str = LIGNE_SOURCE) -> None:
    bloc(d, M, PIED_Y, 8.4, 0.24, [T(texte, taille=8.5, couleur=DISCRET)])
    bloc(d, DROITE - 3.8, PIED_Y, 3.8, 0.24, [T(PIED_IDENTITE, taille=8.5, couleur=DISCRET)], align=PP_ALIGN.RIGHT)


# --------------------------------------------------------------------------------------------------------------------
# Graphiques — matplotlib, dans la palette du dashboard
# --------------------------------------------------------------------------------------------------------------------

_familles = {f.name for f in font_manager.fontManager.ttflist}
plt.rcParams.update({
    "font.family": "Liberation Sans" if "Liberation Sans" in _familles else "DejaVu Sans",
    "font.size": 8.5, "axes.edgecolor": FILET, "axes.labelcolor": ENCRE2, "xtick.color": ENCRE2,
    "ytick.color": ENCRE2, "axes.spines.top": False, "axes.spines.right": False, "figure.facecolor": "white",
    "axes.facecolor": "white", "axes.titlesize": 8.5, "axes.titleweight": "bold", "axes.titlecolor": ENCRE,
})


def _enregistrer(figure, nom: str) -> Path:
    chemin = ASSETS / f"{nom}.png"
    figure.savefig(chemin, dpi=200, bbox_inches="tight", facecolor="white")
    plt.close(figure)
    return chemin


def _grille(axes, axe="y"):
    axes.grid(axis=axe, color="#efece4", linewidth=0.7)
    axes.set_axisbelow(True)


def _nom_region(nom: str) -> str:
    return nom.replace("Maritime hors Grand Lomé", "Maritime hors\nGrand Lomé")


def graphique_internet(w=7.0, h=2.3) -> Path:
    """Objectif 1 : la courbe et ses classes ; l'accès déclaré par région."""
    figure, (a1, a2) = plt.subplots(1, 2, figsize=(w, h), gridspec_kw={"width_ratios": [1.55, 1], "wspace": 0.62})
    a1.plot(USAGE.index, USAGE.pct_population, color=NAVY, lw=1.8, zorder=2)
    teinte = {"accélération": VERT_SERIE, "ralentissement": ROUGE_SERIE}
    for annee, ligne in USAGE.iterrows():
        couleur = teinte.get(ligne.classe, NAVY)
        # Cercle vide : la classe change avec la période de référence (05, C2 ; 07, O1-02) — dite, pas cachée.
        fragile = pd.notna(ligne.classe_variante_2015) and ligne.classe_variante_2015 != ligne.classe
        taille = 38 if ligne.classe in teinte else 12
        a1.scatter(annee, ligne.pct_population, s=taille, zorder=3, linewidths=1.4, edgecolors=couleur,
                   facecolors="white" if fragile else couleur)
    a1.axhline(SEUIL_USAGE, color=DISCRET, ls="--", lw=0.9)
    a1.text(USAGE.index.min(), SEUIL_USAGE + 1.2, f"seuil de {SEUIL_USAGE} %", fontsize=7.5, color=DISCRET)
    dernier = USAGE.loc[A_U, "pct_population"]
    a1.annotate(f"{fr(dernier, 1)} %", (A_U, dernier), xytext=(-6, 7), textcoords="offset points", ha="right",
                fontsize=8.5, weight="bold", color=NAVY)
    a1.set_ylim(0, 47)
    a1.yaxis.set_major_formatter(lambda v, _: f"{v:.0f} %")
    a1.set_xticks([a for a in USAGE.index if a % 2 == 0])
    a1.set_title("Part de la population utilisant Internet", loc="left")
    _grille(a1)
    poignees = [Line2D([0], [0], marker="o", ls="", markersize=6, markerfacecolor=VERT_SERIE,
                       markeredgecolor=VERT_SERIE, label="accélération"),
                Line2D([0], [0], marker="o", ls="", markersize=6, markerfacecolor=ROUGE_SERIE,
                       markeredgecolor=ROUGE_SERIE, label="ralentissement"),
                Line2D([0], [0], marker="o", ls="", markersize=6, markerfacecolor="white", markeredgecolor=DISCRET,
                       label="cercle vide : classe qui dépend\nde la période de référence")]
    a1.legend(handles=poignees, loc="upper left", bbox_to_anchor=(0.0, 0.80), fontsize=7, frameon=False,
              handletextpad=0.3)

    vue = ACCES.sort_values("estimation_pct_2021_22")
    y = np.arange(len(vue))
    a2.barh(y, vue.estimation_pct_2021_22, color=[COULEUR_REGION[n] for n in vue.index], height=0.62)
    for i, (nom, v) in enumerate(vue.estimation_pct_2021_22.items()):
        a2.text(v + 1.5, i, f"{fr(v, 1)} %", va="center", fontsize=7.5, color=ENCRE, weight="bold")
    a2.set_yticks(y, [_nom_region(n) for n in vue.index], fontsize=7.5)
    for etiquette, nom in zip(a2.get_yticklabels(), vue.index):
        etiquette.set_color(COULEUR_REGION[nom])
        etiquette.set_fontweight("bold")
    a2.set_xlim(0, 85)
    a2.set_xticks([])
    a2.spines["bottom"].set_visible(False)
    a2.set_title("Accès déclaré à Internet,\n15 ans et plus (2021/22)", loc="left")
    return _enregistrer(figure, "internet")


def graphique_marche(w=7.0, h=2.3) -> Path:
    """Objectif 2 : part du chiffre d'affaires investie ; sites radio ajoutés chaque année."""
    figure, (a1, a2) = plt.subplots(1, 2, figsize=(w, h), gridspec_kw={"width_ratios": [1.7, 1], "wspace": 0.28})
    annees = list(INV.index)
    couleurs = [OCRE if a == A_I else NAVY for a in annees]
    a1.bar(annees, INV.taux_investissement_pct, color=couleurs, width=0.62)
    for a, v in INV.taux_investissement_pct.items():
        a1.text(a, v + 0.8, fr(v, 1), ha="center", fontsize=7.5, color=ENCRE, weight="bold")
    a1.axhline(SEUIL_INVESTISSEMENT, color=ROUGE_SERIE, ls="--", lw=1)
    a1.set_xlim(annees[0] - 0.6, annees[-1] + 2.3)
    a1.text(annees[-1] + 0.5, SEUIL_INVESTISSEMENT + 0.8, f"seuil de sous-\ninvestissement :\n{SEUIL_INVESTISSEMENT} %",
            fontsize=6.5, color=ROUGE_SERIE, va="bottom")
    a1.set_ylim(0, 42)
    a1.set_yticks([])
    a1.spines["left"].set_visible(False)
    a1.set_xticks(annees)
    a1.tick_params(axis="x", labelsize=7.5)
    a1.set_title("Part du chiffre d’affaires investie (%)", loc="left")

    ans = [int(a) for a in SITES_S.index]
    a2.bar(ans, SITES_S.values, color=[ROUGE_SERIE if a == ans[-1] else "#3987e5" for a in ans], width=0.6)
    for a, v in zip(ans, SITES_S.values):
        a2.text(a, v + 5, f"+{v:.0f}", ha="center", fontsize=7.5, color=ENCRE, weight="bold")
    a2.axhline(REPERE_SITES, color=DISCRET, ls="--", lw=0.9)
    a2.text(ans[-1] + 0.45, REPERE_SITES + 6, f"repère : {REPERE_SITES}", fontsize=7, color=DISCRET, ha="right")
    a2.set_ylim(0, max(SITES_S.values) * 1.2)
    a2.set_yticks([])
    a2.spines["left"].set_visible(False)
    a2.set_xticks(ans)
    a2.tick_params(axis="x", labelsize=7.5)
    a2.set_title("Sites radio ajoutés dans l’année", loc="left")
    return _enregistrer(figure, "marche")


#: Couleurs des trois séries de la diapositive 5, et leurs pas lisibles pour écrire les chiffres du constat.
SERIE_POP, SERIE_AGENCES, SERIE_MM = BLEUS[1], NAVY, ORANGE_SERIE
TEXTE_POP, TEXTE_MM = BLEUS[3], TEXTE_CATEGORIELLE[ORANGE_SERIE]


def graphique_offre(w=7.0, h=2.3) -> Path:
    """Objectif 3 : part de la population, des agences financières et des points mobile money, par milieu."""
    ordre = ["grand_lome", "autres_villes", "rural"]
    noms = ["Grand Lomé", "Autres villes", "Communes rurales"]
    figure, a = plt.subplots(figsize=(w, h))
    x = np.arange(len(ordre))
    series = [("pop_totale", "Part de la population", SERIE_POP, -0.27),
              ("n_formels", "Part des agences financières", SERIE_AGENCES, 0.0),
              ("n_mm", "Part des points mobile money", SERIE_MM, 0.27)]
    for colonne, libelle, couleur, decalage in series:
        valeurs = PAR_MILIEU.loc[ordre, colonne].values
        a.bar(x + decalage, valeurs, width=0.26, color=couleur, label=libelle)
        for xi, v in zip(x + decalage, valeurs):
            a.text(xi, v + 1.2, f"{fr(v, 1)} %", ha="center", fontsize=7.5, color=ENCRE, weight="bold")
    a.set_xticks(x, noms, fontsize=8.5)
    a.set_ylim(0, 70)
    a.set_yticks([])
    a.spines["left"].set_visible(False)
    a.legend(loc="upper left", fontsize=7.5, frameon=False, ncol=3)
    return _enregistrer(figure, "offre")


#: Libellés en clair des statuts d'accès financier (dictionnaire du dashboard, `_VALEURS`).
STATUTS = ["mobile money uniquement", "mobile money dominant", "desserte faible", "desserte diversifiée"]
LIB_STATUT = {"mobile money uniquement": "mobile money\nseul", "mobile money dominant": "mobile money\ndominant",
              "desserte faible": "desserte\nfaible", "desserte diversifiée": "desserte\ndiversifiée"}


def _geo(niveau: str):
    return gpd.read_file(GEO / f"{niveau}.geojson")[["code", "nom", "geometry"]]


def carte_communes(w=1.9, h=4.6) -> Path:
    """Objectif 4 : les 117 communes par statut d'accès financier."""
    geo = _geo("communes").merge(C[["code", "statut_O4_05"]], on="code")
    figure, a = plt.subplots(figsize=(w, h))
    geo.plot(ax=a, color=geo.statut_O4_05.map(STATUT_O4_05), edgecolor="white", linewidth=0.35)
    a.set_axis_off()
    comptes = C.statut_O4_05.value_counts()
    poignees = [Patch(facecolor=STATUT_O4_05[s], label=f"{LIB_STATUT[s]} ({comptes.get(s, 0)})") for s in STATUTS]
    a.legend(handles=poignees, loc="center left", bbox_to_anchor=(0.92, 0.5), fontsize=7.5, frameon=False,
             title="Communes", title_fontsize=8, alignment="left", labelspacing=0.9)
    return _enregistrer(figure, "carte_communes")


#: Les dix préfectures du diagnostic, dans l'ordre du classement (classe, puis population) : numérotées sur la carte.
DIX = pd.concat([P1.sort_values("pop_totale", ascending=False), NC.sort_values("pop_totale", ascending=False)])
COULEUR_CLASSE = {"haute": PRIORITE["haute"], "moyenne": PRIORITE["moyenne"], "faible": PRIORITE["faible"],
                  "non classée": PRIORITE["non classée"]}


def carte_priorites(w=1.9, h=4.6) -> Path:
    """Les 39 préfectures par priorité ; en ocre, les communes où le mobile money est seul (carte de la Synthèse)."""
    geo_p = _geo("prefectures").merge(P[["code", "priorite"]], on="code")
    geo_c = _geo("communes")
    figure, a = plt.subplots(figsize=(w, h))
    geo_p.plot(ax=a, color=geo_p.priorite.map(COULEUR_CLASSE), edgecolor="white", linewidth=0.5)
    geo_c[geo_c.code.isin(MMU.code)].plot(ax=a, color=OCRE, edgecolor="white", linewidth=0.3, alpha=0.9)
    geo_p.dissolve().boundary.plot(ax=a, color="#9a978f", linewidth=0.5)
    points = geo_p.set_index("nom").geometry.representative_point()
    for i, nom in enumerate(DIX.nom, start=1):
        pt = points[nom]
        a.scatter(pt.x, pt.y, s=78, color="white", edgecolor=NAVY, linewidth=0.9, zorder=4)
        a.text(pt.x, pt.y, str(i), ha="center", va="center", fontsize=6, weight="bold", color=NAVY, zorder=5)
    a.set_axis_off()
    comptes = P.priorite.value_counts()
    poignees = [Patch(facecolor=COULEUR_CLASSE[k], edgecolor="#9a978f" if k == "faible" else "none", linewidth=0.5,
                      label=f"{k} ({comptes.get(k, 0)})") for k in ["haute", "moyenne", "faible", "non classée"]]
    poignees.append(Patch(facecolor=OCRE, label=f"commune où\nle mobile\nmoney est\nseul ({len(MMU)})"))
    a.legend(handles=poignees, loc="center left", bbox_to_anchor=(0.92, 0.5), fontsize=7.5, frameon=False,
             title="Priorité", title_fontsize=8, alignment="left", labelspacing=0.8)
    return _enregistrer(figure, "carte_priorites")


def graphique_arbitrage(w=6.0, h=2.55) -> Path:
    """Score de priorité et habitants : l'intensité du manque et son volume ne vont pas ensemble."""
    figure, a = plt.subplots(figsize=(w, h))
    a.axvspan(SEUIL_PRIORITE_HAUTE, 102, color="#e8f0fb", zorder=0)
    a.axvline(SEUIL_PRIORITE_HAUTE, color=NAVY, ls="--", lw=0.9)
    teinte = {"haute": PRIORITE["haute"], "moyenne": PRIORITE["moyenne"], "faible": BLEUS[1]}
    for k in ["haute", "moyenne", "faible"]:
        g = CLASSEES[CLASSEES.priorite == k]
        a.scatter(g.score, g.pop_totale / 1000, s=34, color=teinte[k], edgecolor="white", lw=0.6, zorder=3,
                  label=f"priorité {k} ({len(g)})")
    a.scatter(NC.sans_couv_score, NC.pop_totale / 1000, s=34, facecolor="white", edgecolor="#8a8780", lw=1.2,
              zorder=3, label=f"non classée ({len(NC)}), score sur 2 dimensions")
    for _, ligne in TOP3_POP.iterrows():
        a.annotate(f"{ligne.nom} ({fr(ligne.pop_totale / 1e6, 2)} M)", (ligne.score, ligne.pop_totale / 1000),
                   xytext=(7, -3), textcoords="offset points", fontsize=7.5, weight="bold", color=ENCRE)
    a.text(SEUIL_PRIORITE_HAUTE + 1.5, TOP3_POP.pop_totale.max() / 1000 * 0.97,
           f"priorité haute :\n{SEUIL_PRIORITE_HAUTE} ou plus", fontsize=7.5, color=NAVY, weight="bold", va="top")
    a.set_xlim(0, 102)
    a.set_ylim(0, TOP3_POP.pop_totale.max() / 1000 * 1.08)
    a.set_xlabel("Score de priorité (100 = la préfecture la plus mal servie)", fontsize=8)
    a.set_ylabel("Habitants (milliers)", fontsize=8)
    a.yaxis.set_major_formatter(lambda v, _: fr(v))
    _grille(a, "both")
    a.legend(loc="upper left", bbox_to_anchor=(0.02, 0.58), fontsize=7, frameon=False)
    return _enregistrer(figure, "arbitrage")


def rasteriser_armoiries() -> Path:
    """`python-pptx` ne lit pas le SVG : les armoiries du dashboard sont rasterisées une fois."""
    sortie = ASSETS / "armoiries.png"
    subprocess.run(["rsvg-convert", "-a", "-w", "420", str(RACINE / "dashboard" / "static" / "armoiries-togo-ecu.svg"),
                    "-o", str(sortie)], check=True)
    return sortie


ARMOIRIES: Path = ASSETS / "armoiries.png"

# --------------------------------------------------------------------------------------------------------------------
# Les dix diapositives
# --------------------------------------------------------------------------------------------------------------------


def diapo_couverture(prs) -> None:
    """Couverture = synthèse exécutive (relecture « décideur » du défi 1, correction 1) : le résultat avant la méthode."""
    d = nouvelle_diapositive(prs)
    rect(d, 0, 0, SW, SH, fond=NAVY)

    carte(d, M, 0.40, 0.94, 0.94, bord=OR)
    image_ajustee(d, ARMOIRIES, M + 0.09, 0.49, 0.76, 0.76)
    haut, _, bas = MINISTERE.partition(" et de la ")
    lignes = [T(haut.upper(), taille=10.5, gras=True, couleur=BLANC)]
    if bas:
        lignes.append(T(f"ET DE LA {bas.upper()}", taille=10.5, gras=True, couleur=BLANC, avant=2))
    bloc(d, M + 1.14, 0.46, 6.0, 0.82, lignes, ancre=MSO_ANCHOR.MIDDLE)
    bloc(d, 7.3, 0.50, DROITE - 7.3, 0.40, [T(MARQUE, taille=19, gras=True, couleur=JAUNE, police=POLICE_TITRE)],
         align=PP_ALIGN.RIGHT)
    bloc(d, 7.3, 0.94, DROITE - 7.3, 0.28, [T(BANDEAU, taille=10.5, couleur=BLEU_CLAIR_TEXTE)], align=PP_ALIGN.RIGHT)

    bloc(d, M, 1.66, CW, 0.62, [T("Numérique et inclusion financière au Togo", taille=34, gras=True, couleur=BLANC,
                                   police=POLICE_TITRE)])
    bloc(d, M, 2.30, CW, 0.34, [T("Où agir en priorité pour accélérer l’usage d’Internet et étendre l’accès aux "
                                  "services financiers numériques ?", taille=15, couleur=BLEU_CLAIR_TEXTE)])

    y, h, gap = 2.86, 1.60, 0.22
    w = (CW - 3 * gap) / 4
    hors = "toutes hors du Maritime et du Grand Lomé" if HORS_MARITIME else ""
    cartes = [
        dict(surtitre_txt=f"Internet ({N['usage_annee']})", valeur=pct(N["usage"]),
             phrase="de la population utilise Internet",
             note=f"seuil de {SEUIL_USAGE} % ; × {nb(MULT_2017, 1)} depuis 2017",
             etiquette="Sous le seuil" if N["usage"] < SEUIL_USAGE else "Seuil franchi",
             ton="alerte" if N["usage"] < SEUIL_USAGE else "ok"),
        dict(surtitre_txt="Prix de la data", valeur=pct(N["cout_1go"], 2),
             phrase="du revenu mensuel pour 1 Go", note=f"en {N['cout_annee']} ; seuil d’accessibilité : 2 %",
             etiquette="Non abordable", ton="alerte"),
        dict(surtitre_txt="Accès financier", valeur=str(len(MMU)), unite="communes",
             phrase="sans aucune agence financière",
             note=f"{nb(MMU.pop_totale.sum())} habitants ; le mobile money y est seul",
             etiquette="Priorité absolue", ton="critique", couleur_valeur=TONS["rouge"]["encre"]),
        dict(surtitre_txt="Territoires", valeur=str(len(P1)), unite="préfectures", phrase="en priorité haute",
             note=f"{nb(P1.pop_totale.sum())} habitants{', ' + hors if hors else ''}",
             etiquette="Priorité haute", ton="ok", couleur_valeur=NAVY),
    ]
    for i, c in enumerate(cartes):
        cartouche(d, M + i * (w + gap), y, w, h, **c)

    tous = f"dans les {COMMUNES_MM} communes" if COMMUNES_MM == len(C) else f"dans {COMMUNES_MM} communes"
    banniere(d, M, 4.66, CW, 0.62, "Constat",
             f"Le mobile money est déjà présent {tous} ; l’agence financière manque dans {len(MMU)}, la couverture "
             f"réseau reste à mesurer et la data est chère. Le déficit est rural"
             + (" et se situe hors du Maritime." if HORS_MARITIME else "."), "constat", taille=11.5)
    banniere(d, M, 5.40, CW, 0.62, "Dans l’année",
             f"une première agence financière dans les {len(MMU)} communes où le mobile money est seul, et la mesure "
             f"de la couverture réelle dans les {len(R4A)} communes où elle est inconnue ou douteuse.", "limite",
             taille=11.5)

    surtitre(d, M, 6.24, 4.5, "Présenté par", couleur=JAUNE)
    bloc(d, M, 6.48, 4.5, 0.36, [T(PRESENTE_PAR, taille=16, gras=True, couleur=BLANC)])
    if URL_DASHBOARD:
        surtitre(d, 5.2, 6.24, DROITE - 5.2, "Tableau de bord interactif", couleur=JAUNE)
        lien(d, 5.2, 6.50, DROITE - 5.2, URL_DASHBOARD, URL_DASHBOARD, taille=12, couleur=BLEU_CLAIR_TEXTE)
    mois = ["janvier", "février", "mars", "avril", "mai", "juin", "juillet", "août", "septembre", "octobre",
            "novembre", "décembre"]
    aujourd_hui = datetime.date.today()
    bloc(d, M, PIED_Y, 8.4, 0.24, [T(PIED_IDENTITE, taille=8.5, couleur="#9fb2c6")])
    bloc(d, DROITE - 3.0, PIED_Y, 3.0, 0.24, [T(f"{mois[aujourd_hui.month - 1].capitalize()} {aujourd_hui.year}",
                                                 taille=8.5, couleur="#9fb2c6")], align=PP_ALIGN.RIGHT)


def diapo_methode(prs) -> None:
    """La seule diapositive de méthode, volontairement courte (relecture du défi 1, correction 2)."""
    d = nouvelle_diapositive(prs)
    y = en_tete(d, 2, "01 · Cadrage et méthode", "Comment ces résultats ont-ils été produits ?")

    banniere(d, M, y, CW, 0.60, "Question publique",
             "Où agir en priorité pour accélérer l’usage d’Internet et étendre l’accès aux services financiers "
             "numériques ?", "sombre", taille=12.5)

    surtitre(d, M, y + 0.80, CW, "La chaîne de traitement")
    etapes = [
        ("Données publiques", "recensement de 2022, agences, mobile money, marché des télécoms"),
        ("Diagnostic", f"les 5 objectifs, lus dans {len(C)} communes, {len(P)} préfectures et {len(REGIONS)} régions"),
        ("Priorisation", f"un classement des {len(P)} préfectures, de 0 à 100, puis la nature du manque"),
        ("Recommandations", f"{len(SR)} actions, chacune chiffrée en habitants, avec son horizon"),
    ]
    gap = 0.36
    w = (CW - 3 * gap) / 4
    for i, (titre, detail) in enumerate(etapes):
        x = M + i * (w + gap)
        carte(d, x, y + 1.10, w, 1.06)
        carre_numero(d, x + 0.18, y + 1.26, 0.34, str(i + 1), NAVY, taille=12)
        bloc(d, x + 0.62, y + 1.23, w - 0.78, 0.40, [T(titre, taille=13, gras=True, couleur=NAVY, police=POLICE_TITRE)],
             ancre=MSO_ANCHOR.MIDDLE)
        bloc(d, x + 0.18, y + 1.66, w - 0.36, 0.44, [T(detail, taille=9.5, couleur=ENCRE2, interligne=1.05)])
        if i < 3:
            bloc(d, x + w + 0.02, y + 1.40, gap - 0.04, 0.34, [T("→", taille=16, gras=True, couleur=GRIS_POINT)],
                 align=PP_ALIGN.CENTER)

    surtitre(d, M, y + 2.38, CW, f"Trois dimensions pour classer les {len(P)} préfectures")
    dimensions = [("acces", "Accès aux agences financières", "habitants par agence de banque, de microfinance ou "
                                                             "d’assurance"),
                  ("maillage", "Maillage mobile money", "habitants par point mobile money"),
                  ("couverture", "Couverture réseau", "part de la population couverte, estimation théorique")]
    gap = 0.26
    w = (CW - 2 * gap) / 3
    for i, (cle, titre, detail) in enumerate(dimensions):
        x = M + i * (w + gap)
        carte(d, x, y + 2.68, w, 0.92)
        rect(d, x + 0.20, y + 2.86, 0.26, 0.26, fond=DIM_POINT[cle], rayon=0.07)
        bloc(d, x + 0.60, y + 2.80, w - 0.78, 0.34, [T(titre, taille=13, gras=True, couleur=DIM_TEXTE[cle])],
             ancre=MSO_ANCHOR.MIDDLE)
        bloc(d, x + 0.60, y + 3.16, w - 0.78, 0.36, [T(detail, taille=10, couleur=ENCRE2)])

    banniere(d, M, y + 3.80, CW, 0.64, "À retenir",
             f"le classement va de 0 à 100 ; priorité haute à {SEUIL_PRIORITE_HAUTE} ou plus. Il mesure l’intensité du "
             "manque, pas le nombre d’habitants, et ne porte pas sur l’usage d’Internet, connu seulement par région.",
             "sombre")
    banniere(d, M, y + 4.56, CW, 0.64, "Réserve",
             f"la couverture réseau est théorique (rayon de 20 km autour des antennes) et inconnue dans {len(NC)} "
             "préfectures : toute action sur le réseau attend une mesure de la couverture réelle.", "limite")
    pied(d, f"Sources, seuils et tests de robustesse : tableau de bord, page « {PAGE_METHODE} ».")


def _gabarit_objectif(d, y, cartes, graphique, titre_graphique, constat, enjeu, message):
    """Diapositive d'objectif : quatre chiffres, un graphique, le constat clé et l'enjeu (correction 3 du défi 1)."""
    gap, h = 0.22, 1.60
    w = (CW - 3 * gap) / 4
    for i, c in enumerate(cartes):
        cartouche(d, M + i * (w + gap), y, w, h, **c)
    haut = y + h + 0.18
    largeur_panneau = 7.35
    bas = 6.20
    image_y = panneau(d, M, haut, largeur_panneau, bas - haut, titre_graphique)
    image_ajustee(d, graphique, M + 0.16, image_y, largeur_panneau - 0.32, bas - image_y - 0.10)

    x = M + largeur_panneau + 0.25
    w = DROITE - x
    h_encadre = (bas - haut - 2 * 0.28 - 0.14) / 2
    surtitre(d, x, haut, w, "Constat clé")
    encadre(d, x, haut + 0.28, w, h_encadre, "bleu", constat, taille=11)
    surtitre(d, x, haut + 0.42 + h_encadre, w, "Enjeu pour l’action publique")
    encadre(d, x, haut + 0.70 + h_encadre, w, h_encadre, "vert", enjeu, taille=11)
    etiquette, texte, style = message
    banniere(d, M, 6.32, CW, 0.58, etiquette, texte, style, taille=10.5)
    pied(d)


def diapo_internet(prs, graphiques) -> None:
    d = nouvelle_diapositive(prs)
    y = en_tete(d, 3, "02 · Usage d’Internet", "L’usage d’Internet a triplé depuis 2017, mais il ralentit")
    var_n = USAGE.loc[A_U, "variation_points"]
    cartes = [
        dict(surtitre_txt=f"Usage ({A_U})", valeur=pct(N["usage"]), phrase="de la population utilise Internet",
             note=f"seuil de {SEUIL_USAGE} % ; Afrique subsaharienne : {pct(N['usage_ass'])}",
             etiquette="Sous le seuil" if N["usage"] < SEUIL_USAGE else "Seuil franchi",
             ton="alerte" if N["usage"] < SEUIL_USAGE else "ok"),
        dict(surtitre_txt="Depuis 2017", valeur=f"× {nb(MULT_2017, 1)}",
             phrase=f"de {pct(D1.loc[2017, 'pct_population'])} à {pct(D1.loc[A_U, 'pct_population'])}",
             note=f"deux poussées : {et([str(a) for a in ANNEES_ACCEL])}", couleur_valeur=NAVY),
        dict(surtitre_txt="Rythme récent", valeur=f"+{nb(var_n, 1)}", unite="point",
             phrase=f"gagné en {A_U}",
             note=f"+{nb(VAR_PIC, 1)} points en {max(ANNEES_ACCEL)} ; moins de 2 par an depuis {DEPUIS_LENT}",
             etiquette="Ralentit", ton="alerte"),
        dict(surtitre_txt="Écart régional (2021/22)", valeur=f"× {nb(ACCES_GL / ACCES_SAV, 1)}",
             phrase="d’accès déclaré, Grand Lomé face aux Savanes",
             note=f"{pct(ACCES_GL)} contre {pct(ACCES_SAV)} des 15 ans et plus", couleur_valeur=NAVY),
    ]
    constat = [T(f"Au rythme actuel, l’usage généralisé (60 %) n’arriverait qu’en {int(TENDANCIEL.annee_60pct)} ; "
                 f"au rythme de 2019-2024, dès {int(ACCELERE.annee_60pct)}.")]
    enjeu = [T("Lever les freins de prix et de compétences, d’abord dans les "), R("Savanes"),
             T(f" ({pct(ACCES_SAV)}) et les "), R("Plateaux"), T(f" ({pct(ACCES_PLA)}), où l’accès est le plus bas.")]
    message = ("À savoir", "l’usage national est une estimation internationale ; l’accès régional est déclaré par les "
               "ménages : les deux ne se comparent pas. Un cercle vide marque une classe qui dépend de la période de "
               "référence.", "limite")
    _gabarit_objectif(d, y, cartes, graphiques["internet"], "L’usage d’Internet depuis 2010, et l’accès par région",
                      constat, enjeu, message)


def diapo_marche(prs, graphiques) -> None:
    d = nouvelle_diapositive(prs)
    y = en_tete(d, 4, "03 · Marché des télécoms", "Un duopole qui investit moins, une data encore chère")
    t = lire("07_indicateurs", "o1_04_technologies").set_index("annee")
    ajouts_n, ajouts_0 = SITES_S.iloc[-1], SITES_S.iloc[0]
    cartes = [
        dict(surtitre_txt=f"Marché ({N['marche_annee']})", valeur=pct(N["togocom_data"]),
             phrase="des abonnés data chez Togocom", note="deux opérateurs se partagent le marché",
             etiquette="Très concentré", ton="alerte"),
        dict(surtitre_txt=f"Haut débit ({N['haut_debit_annee']})", valeur=pct(N["haut_debit"]),
             phrase="des abonnements data en 3G ou 4G",
             note=f"un abonnement n’est pas une personne ; {nb(t.loc[N['haut_debit_annee'], 'abonnes_data_mobile'] / 1e6, 1)}"
                  " million d’abonnements",
             etiquette="Seuil franchi" if str(N["haut_debit_statut"]).startswith("acquise") else "Sous le seuil",
             ton="ok" if str(N["haut_debit_statut"]).startswith("acquise") else "alerte"),
        dict(surtitre_txt="Investissement", valeur=pct(INV.loc[A_I, "taux_investissement_pct"]),
             phrase=f"du chiffre d’affaires investi en {A_I}",
             note=f"{pct(INV.loc[A_I0, 'taux_investissement_pct'])} en {A_I0} ; seuil : {SEUIL_INVESTISSEMENT} %",
             etiquette="Proche du seuil", ton="alerte", couleur_valeur=TEXTE_CATEGORIELLE[OCRE]),
        dict(surtitre_txt=f"Sites radio ({int(SITES_S.index[-1])})", valeur=f"+{ajouts_n:.0f}",
             phrase="sites ajoutés dans l’année",
             note=f"+{ajouts_0:.0f} en {int(SITES_S.index[0])} ; repère : {REPERE_SITES} par an",
             etiquette="Sous le repère" if ajouts_n < REPERE_SITES else "Au-dessus", ton="alerte",
             couleur_valeur=TONS["rouge"]["encre"]),
    ]
    constat = [T("Le haut débit mobile est acquis, mais l’investissement et les nouveaux sites reculent depuis "
                 "2023 : l’extension du réseau vers les zones mal couvertes en dépend.")]
    enjeu = [T(f"Garder l’investissement au-dessus de {SEUIL_INVESTISSEMENT} % du chiffre d’affaires, orienter les "
               f"nouveaux sites vers les zones sans réseau, et ramener 1 Go de {pct(R6.base_2025_pct, 2)} à "
               f"{pct(R6.cible_pct, 0)} du revenu mensuel.")]
    message = ("Veille", f"si la baisse de 2024-2025 se répétait, l’investissement tomberait à "
               f"{pct(R8.valeur_2026_si_la_variation_2024_2025_se_repete)} en 2026, sous le seuil : c’est une alerte, "
               "pas une prévision (la série monte et descend par cycles).", "limite")
    _gabarit_objectif(d, y, cartes, graphiques["marche"], "L’investissement des opérateurs et les nouveaux sites radio",
                      constat, enjeu, message)


def diapo_offre(prs, graphiques) -> None:
    d = nouvelle_diapositive(prs)
    y = en_tete(d, 5, "04 · Offre financière", "Le mobile money est dans chaque commune, l’agence financière non")
    gl = PARTS.loc["Grand Lomé"]
    petit, gros = FRAIS.loc[FRAIS.index.min()], FRAIS.loc[FRAIS.index.max()]
    cartes = [
        dict(surtitre_txt="Mobile money", valeur=f"{COMMUNES_MM} / {len(C)}",
             phrase="communes ont le mobile money",
             note=f"{nb(C.n_mm.sum())} points recensés en 2021/2022", etiquette="Partout", ton="ok",
             couleur_valeur=TEXTE_MM),
        dict(surtitre_txt="Grand Lomé", valeur=pct(gl.part_n_formels), phrase="des agences financières",
             note=f"pour {pct(gl.part_pop_totale)} de la population et {pct(gl.part_superficie_km2)} du territoire",
             etiquette="Concentré", ton="alerte", couleur_valeur=NAVY),
        dict(surtitre_txt="Comptes (2024)", valeur=pct(BF["Taux d'activité (%)"]),
             phrase="des comptes mobile money servent",
             note=f"{pct(DETENTION, 0)} des adultes ont un compte ; actif : une opération en 90 jours",
             etiquette="Usage faible", ton="alerte"),
        dict(surtitre_txt="Frais de retrait", valeur=pct(petit.frais_pct_montant),
             phrase=f"pour retirer {nb(petit.name)} FCFA",
             note=f"{pct(gros.frais_pct_montant)} pour {nb(gros.name)} FCFA ; repère : {pct(R5.repere_pct, 0)}",
             etiquette="Régressif", ton="alerte", couleur_valeur=TONS["rouge"]["encre"]),
    ]
    rural = PAR_MILIEU.loc["rural"]
    constat = [T("Les communes rurales abritent "), T(pct(rural.pop_totale), couleur=TEXTE_POP),
               T(" de la population, mais "), T(pct(rural.n_formels), couleur=NAVY),
               T(" des agences et "), T(pct(rural.n_mm), couleur=TEXTE_MM),
               T(f" des points mobile money. L’écart : le mobile money est partout, l’agence manque dans "
                 f"{len(MMU)} communes.")]
    enjeu = [T("Rapprocher l’agence financière du mobile money, déjà présent partout, et rendre les petites opérations "
               f"moins chères : retirer {nb(petit.name)} FCFA coûte {nb(R5.frais_fcfa)} FCFA.")]
    message = ("À savoir", "recensement de 2021/2022 : des lieux, pas des agents ni des transactions. L’usage du "
               "mobile money n’est connu que pour le pays et les régions, jamais commune par commune.", "limite")
    _gabarit_objectif(d, y, cartes, graphiques["offre"],
                      "Population, agences financières et points mobile money, par milieu",
                      constat, enjeu, message)


def diapo_deficits(prs, graphiques) -> None:
    """Objectif 4 : trois manques, trois actions, sur la même ligne (relecture du défi 1, correction 4)."""
    d = nouvelle_diapositive(prs)
    y = en_tete(d, 6, "05 · Déficits", "Quel problème faut-il résoudre ?",
                sous_titre="Trois manques distincts, qui appellent trois actions distinctes.")
    largeur_carte = 3.55
    image_y = panneau(d, M, y, largeur_carte, 6.90 - y, "Les communes selon leur accès",
                      sous_titre="agences financières et mobile money")
    image_ajustee(d, graphiques["carte_communes"], M + 0.12, image_y, largeur_carte - 0.24, 6.90 - image_y - 0.10)

    x = M + largeur_carte + 0.26
    w = DROITE - x
    gap, h = 0.20, 1.72
    wc = (w - 2 * gap) / 3
    cartes = [
        dict(surtitre_txt="Sans agence", valeur=str(len(MMU)), unite="communes",
             phrase="sans aucune agence financière",
             note=f"{nb(MMU.pop_totale.sum())} habitants" + (", toutes rurales" if (MMU.strate_50 == "rural").all()
                                                             else ""),
             etiquette="Priorité absolue", ton="critique", couleur_valeur=TONS["rouge"]["encre"]),
        dict(surtitre_txt="Agence rare", valeur=str(len(SUPPL)), unite="communes",
             phrase="où le mobile money remplace presque l’agence",
             note=f"plus de 20 points par agence ; {pct(PART_SUPPL)} de la population",
             couleur_valeur="#4a3aa7"),
        dict(surtitre_txt="Réseau faible", valeur=str(len(CRIT)), unite="communes",
             phrase="mobile money seul ou dominant, couverture sous 50 %",
             note=f"{nb(CRIT.pop_totale.sum())} habitants",
             etiquette="À confirmer", ton="neutre"),
    ]
    for i, c in enumerate(cartes):
        cartouche(d, x + i * (wc + gap), y, wc, h, **c)

    surtitre(d, x, y + h + 0.18, w, "Trois manques → trois actions")
    colonnes = [("Manque", 2.30, PP_ALIGN.LEFT), ("Où", 2.55, PP_ALIGN.LEFT), ("Habitants", 1.05, PP_ALIGN.RIGHT),
                ("Action à engager", w - 0.14 - 5.90, PP_ALIGN.LEFT)]
    lignes = [
        [("Accès aux agences financières", DIM_TEXTE["acces"], True), f"{len(R1)} communes sans agence",
         nb(SR.loc["R1", "population"]), ("Ouvrir une première agence dans chacune", ENCRE, True)],
        [("Maillage mobile money", DIM_TEXTE["maillage"], True),
         f"{len(R3C)} communes et {len(R3P)} préfectures au réseau mince", nb(SR.loc["R3", "population"]),
         ("Ajouter des points mobile money", ENCRE, True)],
        [("Couverture réseau", DIM_TEXTE["couverture"], True),
         f"{len(R4A)} communes à mesurer ; {len(R4B_P1)} préfectures sous 85 %", nb(SR.loc["R4a", "population"]),
         ("Mesurer, puis étendre le réseau", ENCRE, True)],
    ]
    bas_tableau = tableau_texte(d, x, y + h + 0.46, w, colonnes, lignes, taille=10.5, h_ligne=0.54)
    banniere(d, x, bas_tableau + 0.16, w, 0.66, "Le trait commun : la distance",
             [T("dans les préfectures en priorité haute, "), T(pct(FACT.mediane_priorite_1, 0), couleur=JAUNE),
              T(" des points mobile money sont à plus de 10 km d’une agence financière, contre "),
              T(pct(FACT.mediane_autres_classees, 0), couleur=JAUNE), T(" ailleurs.")], "sombre", taille=11)
    pied(d)


def _nature(nom: str) -> tuple[str, str]:
    """Nature du manque (09, section 7.3), dérivée des déficits marqués et du moteur de la fiche."""
    fiche = FICHES.loc[nom]
    marques = len(str(fiche.deficits_marques).split(", "))
    if marques == 3:
        return "les trois manques", TONS["rouge"]["encre"]
    if fiche.moteur == "accès formel":
        return "agences financières d’abord", DIM_TEXTE["acces"]
    if fiche.moteur == "maillage mobile money":
        return "réseau mobile money mince", DIM_TEXTE["maillage"]
    fragile = str(fiche.confiance_P13) != "élevée"
    return ("couverture d’abord, à confirmer" if fragile else "couverture d’abord"), DIM_TEXTE["couverture"]


def _confiance(v) -> tuple[str, str]:
    v = str(v)
    if v.startswith("non classée"):
        return "non classée", DISCRET
    return v, {"élevée": TONS["vert"]["encre"], "moyenne": TONS["ambre"]["encre"]}.get(v, TONS["rouge"]["encre"])


def diapo_territoires(prs, graphiques) -> None:
    """Où agir en priorité : la carte et les dix préfectures, avec la nature du manque (correction 5 du défi 1)."""
    d = nouvelle_diapositive(prs)
    lieu = ", toutes hors du Maritime et du Grand Lomé" if HORS_MARITIME else ""
    y = en_tete(d, 7, "06 · Territoires", "Où agir en priorité ?",
                sous_titre=f"{len(P1)} préfectures en priorité haute et {len(NC)} non classées{lieu}.")
    largeur_carte = 3.55
    image_y = panneau(d, M, y, largeur_carte, 6.90 - y, f"Les {len(P)} préfectures par priorité",
                      sous_titre="numéros : les dix préfectures du tableau")
    image_ajustee(d, graphiques["carte_priorites"], M + 0.12, image_y, largeur_carte - 0.24, 6.90 - image_y - 0.10)

    x = M + largeur_carte + 0.26
    w = DROITE - x
    gap, h = 0.20, 1.26
    wc = (w - gap) / 2
    confiances = P1.confiance_P13.value_counts()
    carte(d, x, y, wc, h)
    bloc(d, x + 0.18, y + 0.14, wc - 0.36, 0.22, [T("PRIORITÉ HAUTE", taille=8.5, gras=True, couleur=NAVY)])
    bloc(d, x + 0.18, y + 0.38, wc - 0.36, 0.46,
         [dict(runs=[T(str(len(P1)), taille=25, gras=True, couleur=NAVY, police=POLICE_TITRE),
                     T(f" préfectures · {nb(P1.pop_totale.sum())} habitants", taille=12, gras=True, couleur=NAVY)])],
         ajuster=False)
    bloc(d, x + 0.18, y + 0.86, wc - 0.36, 0.32,
         [T(f"{pct(PART_P1)} de la population classée ; confiance élevée pour {confiances.get('élevée', 0)}, "
            f"moyenne pour {confiances.get('moyenne', 0)}", taille=9.5, couleur=ENCRE2)])
    x2 = x + wc + gap
    carte(d, x2, y, wc, h)
    bloc(d, x2 + 0.18, y + 0.14, wc - 0.36, 0.22, [T("NON CLASSÉES", taille=8.5, gras=True, couleur=NAVY)])
    bloc(d, x2 + 0.18, y + 0.38, wc - 0.36, 0.46,
         [dict(runs=[T(str(len(NC)), taille=25, gras=True, couleur=DISCRET, police=POLICE_TITRE),
                     T(f" préfectures · {et(NC.sort_values('pop_totale', ascending=False).nom)}", taille=12,
                       gras=True, couleur=DISCRET)])], ajuster=False)
    haute_2d = (NC.sans_couv_classe == "priorité 1").all()
    bloc(d, x2 + 0.18, y + 0.86, wc - 0.36, 0.32,
         [T("couverture inconnue" + (" ; en priorité haute sur les deux autres dimensions" if haute_2d else ""),
            taille=9.5, couleur=ENCRE2)])

    colonnes = [("#", 0.36, PP_ALIGN.LEFT), ("Préfecture", 1.55, PP_ALIGN.LEFT), ("Région", 1.95, PP_ALIGN.LEFT),
                ("Habitants", 1.00, PP_ALIGN.RIGHT), ("Nature du manque", 2.40, PP_ALIGN.LEFT),
                ("Confiance", w - 0.14 - 7.26, PP_ALIGN.LEFT)]
    lignes = []
    for i, (_, ligne) in enumerate(DIX.iterrows(), start=1):
        nature, couleur = _nature(ligne.nom)
        niveau, couleur_niveau = _confiance(ligne.confiance_P13)
        lignes.append([(str(i), DISCRET, True), (ligne.nom, ENCRE, True),
                       (ligne.unite_regionale, COULEUR_REGION[ligne.unite_regionale], True),
                       (nb(ligne.pop_totale), ENCRE, False), (nature, couleur, True), (niveau, couleur_niveau, True)])
    bas = tableau_texte(d, x, y + h + 0.18, w, colonnes, lignes, taille=10, h_ligne=0.27)
    bloc(d, x, bas + 0.10, w, 6.90 - bas - 0.10,
         [T("Confiance élevée : la priorité ne change sous aucun des poids testés et ne dépend pas de la couverture. "
            "Moyenne : elle repose en partie sur une couverture douteuse.", taille=9, italique=True, couleur=DISCRET)])
    pied(d)


def diapo_arbitrage(prs, graphiques) -> None:
    """L'arbitrage, et les règles qui le tranchent (relecture du défi 1, correction 6 : ne pas laisser la question
    ouverte). Les deux règles sont celles du projet : classe puis population (02) ; priorité absolue des communes sans
    agence, quelle que soit leur préfecture (P17 du 09)."""
    d = nouvelle_diapositive(prs)
    y = en_tete(d, 8, "07 · Arbitrage", "Priorité n’est pas population",
                sous_titre="Le classement mesure l’intensité du manque ; la population en mesure le volume. "
                           "Les deux divergent.")
    largeur = 6.40
    haut_regles = 5.20
    image_y = panneau(d, M, y, largeur, haut_regles - 0.14 - y, "Score de priorité et habitants, par préfecture")
    image_ajustee(d, graphiques["arbitrage"], M + 0.14, image_y, largeur - 0.28, haut_regles - 0.14 - image_y - 0.08)

    x = M + largeur + 0.25
    w = DROITE - x
    surtitre(d, x, y, w, "Deux constats")
    classes_top = {P.set_index("nom").priorite[n] for n in TOP3_POP.nom}
    lecture_top = " ou ".join(k for k in ["moyenne", "faible", "haute"] if k in classes_top)
    constats = [
        ("Les plus peuplées ne sont pas les plus prioritaires",
         [T(f"Les {len(P1)} préfectures en priorité haute comptent {nb(P1.pop_totale.sum())} habitants "
            f"({pct(PART_P1)}). {et(list(TOP3_POP.nom))} ({millions(TOP3_POP.pop_totale.sum())} d’habitants à elles "
            f"trois) sont en priorité {lecture_top}.", taille=10.5, couleur=ENCRE)]),
        ("La préfecture cache des communes",
         [T(f"{len(MMU_AILLEURS)} des {len(R1)} communes sans agence financière ({nb(MMU_AILLEURS.pop_totale.sum())} "
            "habitants) sont dans des préfectures qui ne sont pas en priorité haute : le classement préfectoral ne "
            "les voit pas.", taille=10.5, couleur=ENCRE)]),
    ]
    hc = (haut_regles - 0.14 - (y + 0.30) - 0.14) / 2
    for i, (titre, texte) in enumerate(constats):
        yc = y + 0.30 + i * (hc + 0.14)
        carte(d, x, yc, w, hc)
        carre_numero(d, x + 0.18, yc + 0.16, 0.32, str(i + 1), NAVY, taille=11)
        bloc(d, x + 0.62, yc + 0.14, w - 0.80, 0.36, [T(titre, taille=12, gras=True, couleur=NAVY,
                                                          police=POLICE_TITRE)], ancre=MSO_ANCHOR.MIDDLE)
        bloc(d, x + 0.18, yc + 0.58, w - 0.36, hc - 0.68, [dict(runs=texte, interligne=1.08)])

    surtitre(d, M, haut_regles, CW, "Deux règles pour décider")
    regles = [("rouge", "La commune sans agence passe en premier",
               "priorité absolue, quelle que soit la classe de sa préfecture"),
              ("bleu", "La classe d’abord, la population ensuite",
               "à classe égale, la préfecture la plus peuplée passe devant")]
    wr = (CW - 0.26) / 2
    for i, (ton, titre, texte) in enumerate(regles):
        xr = M + i * (wr + 0.26)
        rect(d, xr, haut_regles + 0.28, wr, 0.76, fond=TONS[ton]["bg"], rayon=0.12, bord=TONS[ton]["bord"])
        bloc(d, xr + 0.22, haut_regles + 0.36, wr - 0.44, 0.30,
             [T(titre, taille=12.5, gras=True, couleur=TONS[ton]["encre"], police=POLICE_TITRE)])
        bloc(d, xr + 0.22, haut_regles + 0.68, wr - 0.44, 0.28, [T(texte, taille=10.5, couleur=ENCRE)])
    banniere(d, M, haut_regles + 1.16, CW, 0.52, "Recommandation",
             "le classement fixe l’urgence, la population dimensionne l’effort, et chaque commune sans agence est "
             "traitée en priorité absolue.", "sombre", taille=11.5)
    pied(d)


def diapo_plan(prs) -> None:
    """LA diapositive centrale : le plan d'action ordonné (relecture du défi 1, correction 7). Chaque effectif est lu
    dans les tables des recommandations ; l'ordre est celui du 10, section 2."""
    d = nouvelle_diapositive(prs)
    y = en_tete(d, 9, "08 · Plan d’action", "Par où commencer ?")
    r4b_cible = int(R4B_P1.pop_a_couvrir_pour_la_cible.sum())
    premieres = ", ".join(R1.nom.head(3))
    entieres = et(MESUREES_ENTIERES) if MESUREES_ENTIERES else ""
    actions = [
        ("rouge", f"Une agence financière dans chacune des {len(R1)} communes où le mobile money est seul",
         f"{len(R1)} agences la première année, {int(R1.points_cible_3ans_tendu.sum())} à 3 ans ; d’abord les plus "
         f"peuplées : {premieres}.",
         [(SR.loc["R1", "horizon"], "ambre"), ("Priorité absolue", "rouge")],
         f"{nb(SR.loc['R1', 'population'])} habitants", "Banques, microfinance, assurances"),
        ("bleu", f"Mesurer la couverture réelle dans {len(R4A)} communes",
         (f"Dont toutes celles de {entieres} : " if entieres else "") + "un préalable à tout investissement dans "
         "le réseau.",
         [(SR.loc["R4a", "horizon"], "ambre"), ("Préalable", "navy")],
         f"{nb(SR.loc['R4a', 'population'])} habitants", "Régulateur des télécoms, opérateurs"),
        ("navy", f"Des agences et des points mobile money dans les {len(R2)} préfectures prioritaires",
         f"{int(R2.points_a_ajouter.sum())} agences, dont {int(R2.dont_points_R1.sum())} comptées dans l’action 1 ; "
         f"{MM_SEUIL} points mobile money pour franchir un seuil de desserte.",
         [(SR.loc["R2", "horizon"], "ambre"), ("Priorité haute", "navy")],
         f"{nb(SR.loc['R2', 'population'])} habitants", "Établissements financiers, opérateurs"),
        ("violet", "Étendre le réseau et la fibre, une fois la couverture mesurée",
         f"{nb(r4b_cible)} habitants de plus dans la couverture ({len(R4B_P1)} préfectures) ; de la fibre dans les "
         f"{len(R9)} préfectures sans fibre recensée.",
         [(SR.loc["R4b", "horizon"], "ambre"), ("Conditionnelle", "violet")],
         f"{nb(SR.loc['R4b', 'population'])} hors couverture", "Opérateurs, société d’infrastructures numériques"),
        ("orange", "Rendre la data et le mobile money abordables, relever les compétences",
         f"1 Go : de {pct(R6.base_2025_pct, 2)} à {pct(R6.cible_pct, 0)} du revenu ; retrait de "
         f"{nb(1000)} FCFA : de {nb(R5.frais_fcfa)} à {nb(R5.frais_cible_fcfa)} FCFA ; "
         f"{et(sorted(R7.region, reverse=True))} d’abord.",
         [(SR.loc["R6", "horizon"], "ambre"), ("National", "gris")],
         "tout le pays", "Régulateur, émetteurs, ministère du numérique"),
    ]
    h, gap = 0.78, 0.09
    for i, (ton, titre, detail, etiquettes, habitants, acteurs) in enumerate(actions):
        yi = y + i * (h + gap)
        carte(d, M, yi, CW, h)
        carre_numero(d, M + 0.18, yi + 0.18, 0.44, str(i + 1), TONS[ton]["encre"], taille=15)
        bloc(d, M + 0.80, yi + 0.08, 3.95, h - 0.16, [T(titre, taille=12, gras=True, couleur=TONS[ton]["encre"],
                                                         interligne=1.02)], ancre=MSO_ANCHOR.MIDDLE)
        bloc(d, M + 4.95, yi + 0.08, 4.15, h - 0.16, [T(detail, taille=10.5, couleur=ENCRE, interligne=1.05)],
             ancre=MSO_ANCHOR.MIDDLE)
        xd = M + 9.30
        for texte, ton_e in etiquettes:
            xd += pastille(d, xd, yi + 0.10, texte, ton_e, taille=8.5, h=0.22) + 0.08
        bloc(d, M + 9.30, yi + 0.36, CW - 9.30 - 0.16, 0.22, [T(habitants, taille=10.5, gras=True, couleur=ENCRE)])
        bloc(d, M + 9.30, yi + 0.56, CW - 9.30 - 0.16, 0.20, [T(acteurs, taille=8.5, couleur=DISCRET)])

    y_veille = y + 5 * (h + gap) + 0.02
    bloc(d, M, y_veille, CW, 0.26,
         [dict(runs=[T("Veilles annuelles : ", taille=10, gras=True, couleur=NAVY),
                     T(f"investissement d’au moins {SEUIL_INVESTISSEMENT} % du chiffre d’affaires "
                       f"({pct(R8.base_2025_pct)} en {A_I}) ; au moins {REPERE_SITES} sites radio ajoutés par an "
                       f"({R10.ajouts_nets_2025:.0f} en {int(SITES_S.index[-1])}).", taille=10, couleur=ENCRE2)])])
    banniere(d, M, y_veille + 0.36, CW, 0.52, "Règle d’ordonnancement",
             "ce qui est certain n’attend pas (agences, points mobile money) ; le réseau attend la mesure de la "
             "couverture réelle.", "sombre", taille=11.5)
    pied(d, f"{len(SR)} recommandations au total ; les habitants ne s’additionnent pas d’une action à l’autre, les "
            "territoires se recoupent.")


def diapo_decisions(prs) -> None:
    """Conclusion : quatre décisions, quatre limites (relecture du défi 1, corrections 8 et 9)."""
    d = nouvelle_diapositive(prs)
    y = en_tete(d, 10, "09 · Conclusion", "Les quatre décisions à engager")
    decisions = [
        ("rouge", "Ouvrir", f"une première agence financière dans les {len(R1)} communes où le mobile money est seul.",
         "dans l’année"),
        ("bleu", "Mesurer", f"la couverture réelle dans {len(R4A)} communes, avant tout investissement dans le réseau.",
         "dans l’année"),
        ("navy", "Densifier", f"agences et points mobile money dans les {len(R2)} préfectures prioritaires"
                              + (", toutes hors du Maritime." if HORS_MARITIME else "."), "à 3 ans"),
        ("orange", "Rendre abordable", f"la data (1 Go sous {pct(R6.cible_pct, 0)} du revenu) et le petit retrait "
                                       f"(sous {pct(R5.repere_pct, 0)} du montant), partout.", "à 5 ans"),
    ]
    gap, h = 0.22, 1.86
    w = (CW - 3 * gap) / 4
    for i, (ton, verbe, texte, quand) in enumerate(decisions):
        x = M + i * (w + gap)
        carte(d, x, y, w, h)
        carre_numero(d, x + 0.20, y + 0.20, 0.40, str(i + 1), TONS[ton]["encre"], taille=14)
        pastille(d, x + w - 0.20, y + 0.28, quand, "ambre", taille=8.5, h=0.24, droite=True)
        bloc(d, x + 0.20, y + 0.70, w - 0.40, 0.40, [T(verbe, taille=18, gras=True, couleur=TONS[ton]["encre"],
                                                        police=POLICE_TITRE)])
        bloc(d, x + 0.20, y + 1.10, w - 0.40, h - 1.20, [T(texte, taille=10.5, couleur=ENCRE, interligne=1.06)])

    yb = y + h + 0.26
    wl = 7.70
    rect(d, M, yb, wl, 2.02, fond=TONS["limite"]["bg"], rayon=0.12, bord=TONS["limite"]["bord"])
    bloc(d, M + 0.24, yb + 0.14, wl - 0.48, 0.24, [T("CE QUE L’ANALYSE NE DIT PAS", taille=10, gras=True,
                                                     couleur=TONS["limite"]["encre"])])
    limites = ["La couverture est théorique : à mesurer avant d’investir dans le réseau.",
               "L’usage d’Internet et du mobile money n’est connu que par région.",
               "Des points de service recensés en 2021/2022 : des lieux, pas des agents.",
               "Aucun coût : des volumes (agences, points, habitants), pas des budgets."]
    bloc(d, M + 0.24, yb + 0.44, wl - 0.48, 1.06,
         [T(f"•  {l}", taille=10.5, couleur=ENCRE, apres=3) for l in limites])
    bloc(d, M + 0.24, yb + 1.56, wl - 0.48, 0.36,
         [T("L’analyse dit où agir et sur quel manque ; elle ne chiffre pas encore le coût des investissements.",
            taille=10.5, gras=True, couleur=TONS["limite"]["encre"])])

    xr = M + wl + 0.25
    wr = DROITE - xr
    carte(d, xr, yb, wr, 2.02)
    bloc(d, xr + 0.22, yb + 0.14, wr - 0.44, 0.24, [T("EXPLORER LE DÉTAIL", taille=10, gras=True, couleur=OCRE_TEXTE)])
    bloc(d, xr + 0.22, yb + 0.42, wr - 0.44, 0.30, [T("Tableau de bord interactif", taille=13, gras=True, couleur=NAVY,
                                                      police=POLICE_TITRE)])
    if URL_DASHBOARD:
        lien(d, xr + 0.22, yb + 0.80, wr - 0.44, URL_DASHBOARD.removeprefix("https://"), URL_DASHBOARD, taille=9.5)
    bloc(d, xr + 0.22, yb + 1.14, wr - 0.44, 0.78,
         [T("Filtrer par région, par priorité et par milieu ; la fiche de chaque préfecture ; en français et en "
            "anglais.", taille=10, couleur=ENCRE2, interligne=1.05)])

    banniere(d, M, yb + 2.24, CW, 0.66, "",
             [T(f"Le mobile money a déjà atteint les {COMMUNES_MM} communes : il reste à y amener l’agence financière, "
                "un réseau mesuré et une data abordable.", taille=14, police=POLICE_TITRE)], "sombre")
    pied(d)


def construire() -> None:
    global ARMOIRIES
    ARMOIRIES = rasteriser_armoiries()
    graphiques = {
        "internet": graphique_internet(),
        "marche": graphique_marche(),
        "offre": graphique_offre(),
        "carte_communes": carte_communes(),
        "carte_priorites": carte_priorites(),
        "arbitrage": graphique_arbitrage(),
    }
    prs = Presentation()
    prs.slide_width, prs.slide_height = LARGEUR, HAUTEUR
    diapo_couverture(prs)
    diapo_methode(prs)
    diapo_internet(prs, graphiques)
    diapo_marche(prs, graphiques)
    diapo_offre(prs, graphiques)
    diapo_deficits(prs, graphiques)
    diapo_territoires(prs, graphiques)
    diapo_arbitrage(prs, graphiques)
    diapo_plan(prs)
    diapo_decisions(prs)
    prs.save(SORTIE)
    print(f"{SORTIE.name} — {len(prs.slides)} diapositives, {SORTIE.stat().st_size / 1024:.0f} ko")
    for avertissement in AVERTISSEMENTS:
        print(f"  ⚠ {avertissement}")


if __name__ == "__main__":
    construire()
